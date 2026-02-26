"""
Professional DOCX report generator for MangoScan diagnoses.
"""
import io
import os
from datetime import datetime


def generate_report(scan, disease_info):
    """
    Generate a professional Word document report for a scan.
    Returns bytes of the .docx file.
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_ALIGN_VERTICAL
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        import copy
    except ImportError:
        raise ImportError("python-docx is required. Run: pip install python-docx")

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    def add_heading(text, level=1, color=None):
        p = doc.add_heading(text, level=level)
        if color:
            for run in p.runs:
                run.font.color.rgb = RGBColor(*color)
        return p

    def add_para(text, bold=False, italic=False, size=None, color=None, align=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
        if align:
            p.alignment = align
        return p

    def add_bullet(text, style='List Bullet'):
        try:
            p = doc.add_paragraph(text, style=style)
        except Exception:
            p = doc.add_paragraph(f"• {text}")
        return p

    def set_cell_bg(cell, hex_color):
        """Set cell background color."""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

    # ─── HEADER ──────────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run('MangoScan — Disease Diagnosis Report')
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(26, 58, 42)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run('AI-Powered Mango Leaf Disease Classification System')
    sub_run.italic = True
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(80, 120, 90)

    doc.add_paragraph()

    # ─── SUMMARY TABLE ───────────────────────────────────────────────────────
    add_heading('Diagnosis Summary', level=1, color=(26, 58, 42))

    table = doc.add_table(rows=6, cols=2)
    table.style = 'Table Grid'
    table.columns[0].width = Cm(5)
    table.columns[1].width = Cm(11)

    rows_data = [
        ('Report ID', f'#{scan.id:06d}'),
        ('Scan Date & Time', scan.created_at.strftime('%B %d, %Y at %I:%M %p')),
        ('Diagnosis', scan.top_disease),
        ('Confidence', f'{scan.top_confidence:.1f}%'),
        ('Disease Type', disease_info.get('type', 'N/A')),
        ('Severity Level', disease_info.get('severity', 'N/A')),
    ]

    header_color = '2D6A4F'
    row_colors = ['F0F7F4', 'FFFFFF']

    for i, (label, value) in enumerate(rows_data):
        row = table.rows[i]
        label_cell = row.cells[0]
        value_cell = row.cells[1]

        # Header cells
        set_cell_bg(label_cell, header_color)
        label_p = label_cell.paragraphs[0]
        label_run = label_p.add_run(label)
        label_run.bold = True
        label_run.font.color.rgb = RGBColor(255, 255, 255)
        label_run.font.size = Pt(10)

        # Value cells
        set_cell_bg(value_cell, row_colors[i % 2])
        val_p = value_cell.paragraphs[0]
        val_run = val_p.add_run(value)
        val_run.font.size = Pt(10)
        if label == 'Diagnosis':
            val_run.bold = True
            val_run.font.size = Pt(12)
        if label == 'Severity Level' and value == 'High':
            val_run.font.color.rgb = RGBColor(200, 30, 30)
            val_run.bold = True
        elif label == 'Severity Level' and value == 'Medium':
            val_run.font.color.rgb = RGBColor(200, 100, 0)

    doc.add_paragraph()

    # ─── PREDICTION SCORES ───────────────────────────────────────────────────
    add_heading('Prediction Confidence Scores', level=1, color=(26, 58, 42))
    add_para('The following table shows confidence scores for all disease classes:', italic=True, size=10)
    doc.add_paragraph()

    preds = scan.all_predictions or []
    if preds:
        pred_table = doc.add_table(rows=len(preds) + 1, cols=3)
        pred_table.style = 'Table Grid'

        # Header row
        headers = ['Rank', 'Disease / Condition', 'Confidence (%)']
        for j, h in enumerate(headers):
            cell = pred_table.rows[0].cells[j]
            set_cell_bg(cell, '2D6A4F')
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)

        for i, pred in enumerate(preds):
            row = pred_table.rows[i + 1]
            is_top = i == 0
            bg = 'D4EDDA' if is_top else ('F0F7F4' if i % 2 == 0 else 'FFFFFF')

            values = [str(i + 1), pred['label'], f"{pred['confidence']:.2f}%"]
            for j, val in enumerate(values):
                cell = row.cells[j]
                set_cell_bg(cell, bg)
                run = cell.paragraphs[0].add_run(val)
                run.font.size = Pt(10)
                if is_top:
                    run.bold = True

    doc.add_paragraph()

    # ─── DISEASE DESCRIPTION ─────────────────────────────────────────────────
    if scan.top_disease != 'Healthy':
        add_heading('About This Disease', level=1, color=(26, 58, 42))
        sci_name = disease_info.get('scientific_name', '')
        if sci_name and sci_name != 'N/A':
            p = doc.add_paragraph()
            p.add_run('Scientific Name: ').bold = True
            italic_run = p.add_run(sci_name)
            italic_run.italic = True

        add_para(disease_info.get('description', ''), size=10)
        doc.add_paragraph()

        # Symptoms
        if disease_info.get('symptoms'):
            add_heading('Symptoms Observed', level=2, color=(45, 106, 79))
            for s in disease_info['symptoms']:
                add_bullet(s)
            doc.add_paragraph()

        # Causes
        if disease_info.get('causes'):
            add_heading('Causes & Conditions', level=2, color=(45, 106, 79))
            for c in disease_info['causes']:
                add_bullet(c)
            doc.add_paragraph()

    # ─── TREATMENT ───────────────────────────────────────────────────────────
    add_heading('Treatment Recommendations', level=1, color=(26, 58, 42))
    for t in disease_info.get('treatment', []):
        add_bullet(t)
    doc.add_paragraph()

    # ─── PREVENTION ──────────────────────────────────────────────────────────
    add_heading('Prevention Measures', level=1, color=(26, 58, 42))
    for p_item in disease_info.get('prevention', []):
        add_bullet(p_item)
    doc.add_paragraph()

    # ─── SPREAD CONTROL ──────────────────────────────────────────────────────
    add_heading('Spread Control Measures', level=1, color=(26, 58, 42))
    for sc in disease_info.get('spread_control', []):
        add_bullet(sc)
    doc.add_paragraph()

    # ─── ORGANIC OPTIONS ─────────────────────────────────────────────────────
    if disease_info.get('organic_options'):
        add_heading('Organic / Eco-Friendly Options', level=1, color=(26, 58, 42))
        for o in disease_info['organic_options']:
            add_bullet(o)
        doc.add_paragraph()

    # ─── FEEDBACK ────────────────────────────────────────────────────────────
    try:
        fb = scan.feedback
        add_heading('User Feedback', level=1, color=(26, 58, 42))
        fb_table = doc.add_table(rows=4, cols=2)
        fb_table.style = 'Table Grid'
        fb_rows = [
            ('Rating', f"{'★' * fb.rating}{'☆' * (5 - fb.rating)} ({fb.rating}/5)"),
            ('Diagnosis Correct?', 'Yes' if fb.correct_diagnosis else ('No' if fb.correct_diagnosis is False else 'Not specified')),
            ('Correct Disease', fb.correct_disease or 'N/A'),
            ('Comment', fb.comment or 'No comment provided'),
        ]
        for i, (label, value) in enumerate(fb_rows):
            set_cell_bg(fb_table.rows[i].cells[0], '2D6A4F')
            r = fb_table.rows[i].cells[0].paragraphs[0].add_run(label)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(10)
            fb_table.rows[i].cells[1].paragraphs[0].add_run(value).font.size = Pt(10)
        doc.add_paragraph()
    except Exception:
        pass

    # ─── DISCLAIMER ──────────────────────────────────────────────────────────
    doc.add_paragraph()
    disc = doc.add_paragraph()
    disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    disc_run = disc.add_run(
        'DISCLAIMER: This report is generated by an AI system for informational purposes only. '
        'Always consult a certified agricultural extension officer or plant pathologist for professional diagnosis and treatment advice.'
    )
    disc_run.italic = True
    disc_run.font.size = Pt(8)
    disc_run.font.color.rgb = RGBColor(120, 120, 120)

    # ─── FOOTER INFO ─────────────────────────────────────────────────────────
    foot_p = doc.add_paragraph()
    foot_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot_run = foot_p.add_run(
        f'Generated by MangoScan  |  {datetime.now().strftime("%B %d, %Y")}  |  AI-Powered Plant Disease Detection'
    )
    foot_run.font.size = Pt(8)
    foot_run.font.color.rgb = RGBColor(100, 140, 110)

    # Save to bytes
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
